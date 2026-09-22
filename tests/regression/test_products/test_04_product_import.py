"""REG-PROD-IMPORT: ikki yangi ТМЦ, preview/cancel va filialga biriktirish."""

import os
from pathlib import Path
from uuid import uuid4

import allure
import pytest
from openpyxl import Workbook, load_workbook
from playwright.sync_api import expect

from tests.regression.test_products.product_cases import product_case
from tests.regression.test_products.test_01_create_product import check_product_card
from tests.shared.authorization import authorization, current_company_code
from utils.base_pages.auto_base_page import AutoBasePage
from utils.data_store import save_data
from utils.helper_utils import query_int_from_url

pytestmark = [pytest.mark.regression, pytest.mark.regression_requires_setup, allure.epic("Regression"), allure.feature("ТМЦ"), allure.story("Excel import")]

IMPORT_HEADERS = (
    "ИД", "Код", "Штрих-код", "Название(*)", "Альтернативное название",
    "Продукция", "Товар", "Сырьё", "Оборудование", "Производитель", "Регион",
    "Ед. изм.(*)", "Тип кейса (1 уровень)", "Кол-во в кейсе (1 уровень)",
    "Группа", "Категория", "Торговая марка", "Вес нетто", "Вес брутто", "Литр",
    "Порядковый номер", "Артикул код", "Статус", "GTIN", "Кол-во единиц GTIN",
    "IKPU", "ТН ВЭД", "Код группы маркированных товаров",
)


def _write_import_file(template_path, destination, products):
    """Shablon schema'sini tekshiradi; mavjud record/IDlarni ko'chirmay XLSX yaratadi."""
    template = load_workbook(template_path, read_only=True, data_only=True)
    try:
        headers = tuple(cell.value for cell in next(template.active.iter_rows(min_row=1, max_row=1)))
        if headers != IMPORT_HEADERS:
            raise AssertionError(f"ТМЦ import shabloni o'zgargan: {headers!r}")
    finally:
        template.close()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ТМЦ"
    sheet.append(IMPORT_HEADERS)
    for product in products:
        values = {
            "Код": product["code"], "Название(*)": product["name"],
            "Альтернативное название": product["short_name"],
            "Продукция": "Y" if product["production"] else "N", "Товар": "Y",
            "Сырьё": "Y" if product["material"] else "N", "Оборудование": "N",
            "Ед. изм.(*)": product["measure"], "Тип кейса (1 уровень)": product["box_short_name"],
            "Кол-во в кейсе (1 уровень)": product["box_qty"],
            "Вес нетто": product["weight_netto"], "Вес брутто": product["weight_brutto"],
            "Литр": product["volume"], "Порядковый номер": product["order_no"],
            "Артикул код": product["article_code"], "Статус": "A",
        }
        sheet.append([values.get(header) for header in IMPORT_HEADERS])
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    workbook.close()
    # Fixture fayli: faqat sarlavha + testcase qatorlari; barcha IDlar bo'sh.
    saved = load_workbook(destination, read_only=True, data_only=True)
    try:
        rows = list(saved.active.values)
        if len(rows) != len(products) + 1 or any(row[0] is not None for row in rows[1:]):
            raise AssertionError("Import fayliga ortiqcha qator yoki mavjud mahsulot IDsi tushdi")
    finally:
        saved.close()


def _prepare_import_preview(page, code, products, suffix):
    """Ikki import case uchun bir xil admin UI, schema va preview tayyorlovi."""
    base = AutoBasePage(page)
    directory = Path("test-results/downloads") / f"regression-products-{code}-{suffix}"
    directory.mkdir(parents=True, exist_ok=True)
    with allure.step("1 - Import uchun admin profiliga kirish"):
        authorization(page, who="admin")
        base.switch_filial(name="Администрирование")
    with allure.step("2 - Admin ТМЦ listida import kodlari mavjud emasligini tekshirish"):
        base.navigate_to(tab="Справочники", name="ТМЦ")
        base.expect_page(heading="ТМЦ")
        for product in products:
            base.grid_controller(search=product["code"])
            base.grid(state="empty")
    with allure.step("3 - ТМЦ import formasini ochish"):
        # Nomsiz dropdown toggle uchun semantic button helperi yo'q.
        page.locator(".btn-group").filter(has=page.get_by_role("button", name="Создать", exact=True)).locator("button.dropdown-toggle").click()
        base.click(name="Импорт", role="link", exact=True)
        base.expect_page(heading="ТМЦ (импорт)")
    with allure.step("4 - Joriy import mappingi va boshlang'ich qatorni tekshirish"):
        base.click(name="Настройка", exact=True)
        base.input(label="Начальная строка", expect_value="2")
        # Mappingda label elementi yo'q: sg-sub-row label/value juftligini cheklaydi.
        mappings = page.locator('.sg-sub-row:has(input[ng-model="item.column_number"])')
        expect(mappings).to_have_count(len(IMPORT_HEADERS))
        for index, header in enumerate(IMPORT_HEADERS, start=1):
            row = mappings.filter(has=page.get_by_text(header, exact=True))
            expect(row).to_have_count(1)
            base.input(locator=row.locator("input"), expect_value=str(index))
    with allure.step("5 - Saytdagi shablondan faqat yangi testcase qatorli XLSX tayyorlash"):
        base.click(name="Импорт", exact=True)
        with page.expect_download() as download_info:
            base.click(name="Шаблон", exact=True)
        template_path = directory / "template.xlsx"
        download_info.value.save_as(template_path)
        upload_path = directory / "products.xlsx"
        _write_import_file(template_path, upload_path, products)
        allure.attach.file(str(upload_path), name="ТМЦ import: faqat yangi testcase mahsulotlari", attachment_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", extension="xlsx")
    with allure.step("6 - Faylni previewga yuklash va ikkala qatorni tekshirish"):
        page.locator('input[type="file"][accept*=".xlsx"]').set_input_files(str(upload_path))
        base.click(name="Загрузить файл", exact=True)
        base.wait_for_loader()
        for product in products:
            base.grid(product["code"], product["name"], product["short_name"], product["measure"], root='b-pg-grid[name="products"]')
        base.grid(state="empty", root='b-pg-grid[name="error_messages"]')
        expect(page.locator('b-pg-grid[name="products"] .tbl-row').filter(has=page.locator(".tbl-cell"))).to_have_count(len(products))


def run_import_products(page, code, *, save_baseline=False):
    """Testcase: 2 ТМЦ importi → to'plamga biriktirish → test userida karta tekshiruvi.

    1–6. Admin importi, yangi XLSX va preview qiymatlarini tekshirish.
    7–8. Importni saqlash, error grid va listda haqiqiy natijani tekshirish.
    9–12. Kerakli to'plamni ochib, har ikki mahsulotni alohida tasdiq bilan biriktirish.
    13–18. Test useriga kirib, o'z filialida list/view/edit qiymatlarini tekshirish.

    Narx va stock alohida keyingi bosqich; card tayyorligini ularga tenglashtirmaydi.
    """
    products = [product_case(code, key) for key in ("UZS-B", "USD-B")]
    base = AutoBasePage(page)
    if save_baseline:
        for product in products:
            save_data(f"regression_product_{product['key'].lower().replace('-', '_')}", None)
    _prepare_import_preview(page, code, products, "baseline")
    with allure.step("7 - Ikki ТМЦni import orqali saqlash"):
        base.click(name="Загрузить и сохранить", exact=True)
        base.wait_for_loader()
        base.grid(state="empty", root='b-pg-grid[name="error_messages"]')
    with allure.step("8 - Importdan listga qaytib saqlangan ikki mahsulotni topish"):
        base.click(name="Закрыть", exact=True)
        base.expect_page(heading="ТМЦ")
        for product in products:
            base.grid_controller(search=product["code"])
            base.grid(product["code"], product["name"], "Активный")
    with allure.step("9 - Kerakli Наборы ТМЦ listini ochish"):
        base.click(name="Наборы ТМЦ", role="link", exact=True)
        base.expect_page(heading="Наборы ТМЦ")
        base.grid_controller(search=f"sector-pw{code}")
    with allure.step("10 - Faqat test to'plamining biriktirish formasini ochish"):
        base.grid(f"sector-pw{code}", click=True)
        base.click(name="Прикрепить", exact=True)
        base.expect_page(heading="Прикрепленные ТМЦ")
    for product in products:
        with allure.step(f"11 - {product['key']}: available mahsulotni biriktirish"):
            base.click(name="Доступные", exact=True)
            base.expect_page(heading="Доступные ТМЦ")
            base.grid_controller(search=product["code"], root="b-grid-controller:visible")
            base.grid(product["code"], product["name"], root='b-grid[name="inventories"]', click=True)
            base.click(name="Прикрепить", exact=True)
            base.confirm_biruni(expected_text=f"Прикрепить ТМЦ {product['name']}?")
            base.wait_for_loader()
            base.grid(state="empty", root='b-grid[name="inventories"]')
        with allure.step(f"12 - {product['key']}: biriktirilganlar ro'yxatida tekshirish"):
            base.click(name="Прикрепленные", exact=True)
            base.expect_page(heading="Прикрепленные ТМЦ")
            base.grid_controller(search=product["code"], root="b-grid-controller:visible")
            base.grid(product["code"], product["name"], root='b-grid[name="inventories"]')
    with allure.step("13 - Test useri bilan operatsion filialga kirish"):
        authorization(page, who="user", code=code)
        base.switch_filial(name=f"filial-pw{code}")
    with allure.step("14 - Test userida ТМЦ listini ochish"):
        base.navigate_to(tab="Справочники", name="ТМЦ")
        base.expect_page(heading="ТМЦ")
    for product in products:
        with allure.step(f"15 - {product['key']}: test userida view va ID tekshiruvi"):
            base.grid_controller(search=product["code"])
            base.grid(product["code"], product["name"], "Активный", click=True)
            base.click(name="Просмотреть", exact=True)
            base.expect_page(heading="ТМЦ (просмотр)", url="inventory_view?product_id=")
            product_id = query_int_from_url(page.url, "product_id")
        with allure.step(f"16 - {product['key']}: Excel qiymatlarini edit kartasi bilan solishtirish"):
            base.click(name="Изменить", exact=True)
            base.expect_page(heading="ТМЦ (изменение)")
            check_product_card(page, product)
        with allure.step(f"17 - {product['key']}: o'zgartirmasdan viewga qaytish"):
            base.click(name="Закрыть", exact=True)
            base.expect_page(heading="ТМЦ (просмотр)")
        with allure.step(f"18 - {product['key']}: listga qaytish"):
            base.click(name="Закрыть", exact=True)
            base.expect_page(heading="ТМЦ")
        if save_baseline:
            save_data(f"regression_product_{product['key'].lower().replace('-', '_')}", {"run_code": str(code), "server": os.environ["COMPANY_URL"].rstrip("/"), "company": current_company_code(), "product_id": product_id, "card_verified": True, "price_verified": False, "stock_verified": False, **product})


def run_cancel_product_import(page, code):
    """Testcase: 2 valid qator previewsi saqlanmasdan yopilsa record yaratilmaydi.

    1–6. Mustaqil noyob ikki mahsulotli XLSXni previewgacha yuklash.
    7. Importni saqlamasdan yopish; ikkala kod ham listda bo'lmasligi kerak.
    """
    suffix = uuid4().hex[:8]
    products = [product_case(code, key, prefix=f"reg-impcancel-{suffix}") for key in ("UZS-B", "USD-B")]
    base = AutoBasePage(page)
    _prepare_import_preview(page, code, products, suffix)
    with allure.step("7 - Import previewni bekor qilish va yashirin saqlanish yo'qligini tekshirish"):
        base.click(name="Закрыть", exact=True)
        base.expect_page(heading="ТМЦ")
        for product in products:
            base.grid_controller(search=product["code"])
            base.grid(state="empty")


@allure.title("REG-PROD-IMPORT: UZS-B va USD-B importi va filialga biriktirish")
def test_import_products(page, code):
    run_import_products(page, code, save_baseline=True)


@allure.title("REG-PROD-IMPORT-CANCEL: Excel previewni saqlamasdan yopish")
def test_cancel_product_import(page, code):
    run_cancel_product_import(page, code)
