"""TMZ-2/3/4/5/6: Excel, import, view va Цены qadamlari bitta test ichida."""

import os
import re
from pathlib import Path

import allure
import pytest
from openpyxl import Workbook, load_workbook
from playwright.sync_api import expect

from tests.shared.authorization import authorization, current_company_code
from utils.base_pages.auto_base_page import AutoBasePage
from utils.data_store import save_data
from utils.helper_utils import query_int_from_url

pytestmark = [pytest.mark.regression, pytest.mark.regression_requires_setup, allure.epic("Regression"), allure.feature("ТМЦ"), allure.story("Excel va Цены")]

PRODUCT_DATA_PATH = Path(__file__).resolve().with_name("data_store")


def run_import_product(page, user_page, code):
    """1. To'plam. 2. Excel/preview/save. 3. Attach/view. 4. UZS/USD narx.

    Admin va user alohida contextlarda; barcha product qadamlari shu testda.
    """
    user_base = AutoBasePage(user_page)
    admin_base = AutoBasePage(page)
    import_headers = ('ИД', 'Код', 'Штрих-код', 'Название(*)', 'Альтернативное название', 'Продукция', 'Товар', 'Сырьё', 'Оборудование', 'Производитель', 'Регион', 'Ед. изм.(*)', 'Тип кейса (1 уровень)', 'Кол-во в кейсе (1 уровень)', 'Группа', 'Категория', 'Торговая марка', 'Вес нетто', 'Вес брутто', 'Литр', 'Порядковый номер', 'Артикул код', 'Статус', 'GTIN', 'Кол-во единиц GTIN', 'IKPU', 'ТН ВЭД', 'Код группы маркированных товаров')

    with allure.step("1 - User filialida ТМЦ listini ochish"):
        user_base.switch_filial(name=f"filial-pw{code}")
        user_base.navigate_to(tab="Справочники", name="ТМЦ")
        user_base.expect_page(heading="ТМЦ", url="inventory_list")

    with allure.step("2 - Yaratish formasidan birinchi ko'ringan to'plamni aniqlash"):
        user_base.click(name="Создать", exact=True)
        user_base.expect_page(heading="ТМЦ (создание)", url="inventory+add")
        sector = user_base.multiselect(label="Наборы ТМЦ", clear=True, select_first=True, return_value=True)[0]

    with allure.step("3 - Mahsulot yaratmasdan listga qaytish"):
        user_base.click(name="Закрыть", exact=True)
        user_base.expect_page(heading="ТМЦ", url="inventory_list")

    products = [
        {
            "number": number, "name": f"TMZ-{number}-{code}", "code": f"tmz-{number}-{code}",
            "measure": "шт", "kind": "Товар", "sector": sector,
            "weight_netto": str(number * 100), "weight_brutto": str(number * 100 + 10),
            "volume": str(number * 100 + 20), "price_uzs": str(number * 5000),
            "price_usd": str(number * 5), "price_type_uzs": f"Price Type UZB-pw{code}",
            "price_type_usd": f"Price Type USA-pw{code}",
        }
        for number in (2, 3, 4, 5, 6)
    ]
    for product in products:
        save_data(f"regression_tmz_{product['number']}", None, path=PRODUCT_DATA_PATH)

    with allure.step("4 - Alohida admin contextida login va Администрирование"):
        authorization(page, who="admin")
        admin_base.switch_filial(name="Администрирование")

    with allure.step("5 - Admin ТМЦ listida beshta kod yo'qligini tekshirish"):
        admin_base.navigate_to(tab="Справочники", name="ТМЦ")
        admin_base.expect_page(heading="ТМЦ", url="inventory_list")
        for product in products:
            admin_base.grid_controller(search=product["code"])
            admin_base.grid(state="empty")

    with allure.step("6 - ТМЦ import formasini ochish"):
        admin_base.click(locator="button.dropdown-toggle", root='.btn-group:has(button:text-is("Создать"))')
        admin_base.click(name="Импорт", role="link", exact=True)
        admin_base.expect_page(heading="ТМЦ (импорт)")

    with allure.step("7 - Import mappingi va boshlang'ich qatorni tekshirish"):
        admin_base.click(name="Настройка", exact=True)
        admin_base.input(label="Начальная строка", expect_value="2")
        mappings = page.locator('.sg-sub-row:has(input[ng-model="item.column_number"])')
        expect(mappings).to_have_count(len(import_headers))
        for index, header in enumerate(import_headers, start=1):
            row = mappings.filter(has=page.get_by_text(header, exact=True))
            expect(row).to_have_count(1)
            admin_base.input(locator=row.locator("input"), expect_value=str(index))
        admin_base.click(name="Импорт", exact=True)

    with allure.step("8 - Saytdan Excel shablonini yuklab olish"):
        directory = Path("test-results/downloads") / f"regression-tmz-{code}"
        directory.mkdir(parents=True, exist_ok=True)
        with page.expect_download() as download_info:
            admin_base.click(name="Шаблон", exact=True)
        template_path = directory / "template.xlsx"
        upload_path = directory / "TMZ-2-3-4-5-6.xlsx"
        download_info.value.save_as(template_path)

    with allure.step("9 - Shablon schema tekshiruvi va faqat beshta yangi qatorli Excel"):
        template = load_workbook(template_path, read_only=True, data_only=True)
        try:
            headers = tuple(cell.value for cell in next(template.active.iter_rows(min_row=1, max_row=1)))
            if headers != import_headers:
                raise AssertionError(f"ТМЦ import shabloni o'zgargan: {headers!r}")
        finally:
            template.close()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "ТМЦ"
        sheet.append(import_headers)
        for product in products:
            values = {"Код": product["code"], "Название(*)": product["name"], "Товар": "Y", "Продукция": "N", "Сырьё": "N", "Оборудование": "N", "Ед. изм.(*)": product["measure"], "Вес нетто": int(product["weight_netto"]), "Вес брутто": int(product["weight_brutto"]), "Литр": int(product["volume"]), "Статус": "A"}
            sheet.append([values.get(header) for header in import_headers])
        workbook.save(upload_path)
        workbook.close()
        saved = load_workbook(upload_path, read_only=True, data_only=True)
        try:
            excel_rows = list(saved.active.values)
            if len(excel_rows) != len(products) + 1 or any(row[0] is not None for row in excel_rows[1:]):
                raise AssertionError("Excelda ortiqcha qator yoki mavjud mahsulot IDsi bor")
        finally:
            saved.close()
        allure.attach.file(str(upload_path), name="TMZ-2, TMZ-3, TMZ-4, TMZ-5, TMZ-6 import", attachment_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", extension="xlsx")

    with allure.step("10 - Excelni previewga yuklash va input qiymatlarini tekshirish"):
        page.locator('input[type="file"][accept*=".xlsx"]').set_input_files(str(upload_path))
        admin_base.click(name="Загрузить файл", exact=True)
        admin_base.confirm_biruni(expected_text="Загрузить файл?")
        admin_base.wait_for_loader()
        admin_base.grid(state="empty", root='b-pg-grid[name="error_messages"]')
        rows = page.locator('b-pg-grid[name="products"] .tbl-row').filter(has=page.locator('input[ng-model="row.code"]'))
        expect(rows).to_have_count(len(products))
        for product in products:
            with allure.step(f"{product['name']}: preview kodi, nomi va birligi"):
                # Lokal grid name bo'yicha qidiradi; qator qiymatlari inputlarda.
                admin_base.input(ng_model="g.searchValue", root='b-pg-controller[name="products"]', value=product["name"]).press("Enter")
                expect(rows).to_have_count(1)
                for model, key in (("row.code", "code"), ("row.name", "name"), ("row.measure_short_name", "measure")):
                    admin_base.input(ng_model=model, root=rows, expect_value=product[key])
        admin_base.input(ng_model="g.searchValue", root='b-pg-controller[name="products"]', value="").press("Enter")
        expect(rows).to_have_count(len(products))

    with allure.step("11 - Importni saqlash, preview tozalanishi va xato yo'qligi"):
        admin_base.click(name="Загрузить и сохранить", exact=True)
        admin_base.confirm_biruni(expected_text="Вы действительно хотите загрузить и сохранить")
        admin_base.wait_for_loader()
        admin_base.expect_page(heading="ТМЦ (импорт)")
        admin_base.grid(state="empty", root='b-pg-grid[name="products"]')
        admin_base.grid(state="empty", root='b-pg-grid[name="error_messages"]')

    with allure.step("12 - Importni yopish va listda beshta mahsulotni tekshirish"):
        admin_base.click(name="Закрыть", exact=True)
        admin_base.expect_page(heading="ТМЦ", url="inventory_list")
        for product in products:
            admin_base.grid_controller(search=product["code"])
            admin_base.grid(product["code"], product["name"], "Активный")

    with allure.step("13 - Наборы ТМЦ listini ochish"):
        admin_base.click(name="Наборы ТМЦ", role="link", exact=True)
        admin_base.expect_page(heading="Наборы ТМЦ")
        admin_base.grid_controller(search=sector)

    with allure.step("14 - Tanlangan to'plamga biriktirish formasini ochish"):
        admin_base.grid(sector, click=True)
        admin_base.click(name="Прикрепить", exact=True)
        admin_base.expect_page(heading="Прикрепленные ТМЦ")

    for product in products:
        with allure.step(f"15 - {product['name']}: Доступные ichidan biriktirish"):
            admin_base.click(name="Доступные", exact=True)
            admin_base.expect_page(heading="Доступные ТМЦ")
            admin_base.grid_controller(search=product["code"], root="b-grid-controller:visible")
            admin_base.grid(product["code"], product["name"], root='b-grid[name="inventories"]', click=True)
            admin_base.click(name="Прикрепить", exact=True)
            admin_base.confirm_biruni(expected_text=f"Прикрепить ТМЦ {product['name']}?")
            admin_base.wait_for_loader()
        with allure.step(f"16 - {product['name']}: biriktirilganini tekshirish"):
            admin_base.click(name="Прикрепленные", exact=True)
            admin_base.expect_page(heading="Прикрепленные ТМЦ")
            admin_base.grid_controller(search=product["code"], root="b-grid-controller:visible")
            admin_base.grid(product["code"], product["name"], root='b-grid[name="inventories"]')

    with allure.step("17 - Oddiy user filialida ТМЦ listini ochish"):
        user_base.switch_filial(name=f"filial-pw{code}")
        user_base.navigate_to(tab="Справочники", name="ТМЦ")
        user_base.expect_page(heading="ТМЦ", url="inventory_list")

    for product in products:
        numeric_values = {key: re.compile(r"^\s*" + r"\s*".join(product[key]) + r"(?:[.,]0+)?\s*$") for key in ("weight_netto", "weight_brutto", "volume")}
        with allure.step(f"18 - {product['name']}: kod va nom bo'yicha alohida qidirish"):
            user_base.grid_controller(search=product["code"])
            user_base.grid(product["code"], product["name"], product["measure"], "Активный")
            user_base.grid_controller(search=product["name"])
            user_base.grid(product["code"], product["name"], product["measure"], "Активный", click=True)
        with allure.step(f"19 - {product['name']}: viewni ochish va ID"):
            user_base.click(name="Просмотреть", exact=True)
            user_base.expect_page(heading="ТМЦ (просмотр)", url="inventory_view?product_id=")
            product["product_id"] = query_int_from_url(user_page.url, "product_id")
        with allure.step(f"20 - {product['name']}: view maydonlarini solishtirish"):
            user_base.text(product["name"], root="#anor393-span-name")
            user_base.form_view(label="Тип ТМЦ", expect_value=product["kind"])
            for label, key in (("Код", "code"), ("Вес нетто:", "weight_netto"), ("Вес брутто:", "weight_brutto"), ("Литр/Дециметр куб:", "volume")):
                user_base.form_view(
                    label=label,
                    expect_value=product[key] if key == "code" else numeric_values[key],
                    root=".b-offcanvas",
                )
        with allure.step(f"21 - {product['name']}: Изменить orqali birlik/to'plamni o'qish"):
            user_base.click(name="Изменить", exact=True)
            user_base.expect_page(heading="ТМЦ (изменение)")
            user_base.b_input(label="Ед. изм.", expect_value=product["measure"])
            user_base.multiselect(label="Наборы ТМЦ", expect_value=product["sector"])
        with allure.step(f"22 - {product['name']}: o'zgartirishsiz viewga qaytish"):
            user_base.click(name="Закрыть", exact=True)
            user_base.expect_page(heading="ТМЦ (просмотр)")
        with allure.step(f"23 - {product['name']}: listga qaytish"):
            user_base.click(name="Закрыть", exact=True)
            user_base.expect_page(heading="ТМЦ", url="inventory_list")

    for type_key, amount_key, currency in (("price_type_uzs", "price_uzs", "Узбекский сум"), ("price_type_usd", "price_usd", "Доллар США")):
        price_type = products[0][type_key]
        with allure.step(f"24 - {currency}: Цены listini ochish"):
            user_base.navigate_to(tab="Справочники", name="Цены")
            user_base.expect_page(heading="Цены", url="price_type_list")
        with allure.step(f"25 - {currency}: narx turini tekshirib Установка цен formasini ochish"):
            user_base.grid_controller(search=price_type)
            user_base.grid(price_type, currency, click=True)
            user_base.click(name="Установка цен", exact=True)
            user_base.expect_page(heading=f"Установка цен: {price_type}", url="product_price?price_type_id=")
        for product in products:
            with allure.step(f"26 - {product['name']}: {currency} narxini kiritish"):
                user_base.input(ng_model="g.searchValue", root='b-pg-controller[name="products"]', value=product["code"]).press("Enter")
                row = user_base.grid(product["code"], product["name"], root='b-pg-grid[name="products"]')
                user_base.input(ng_model="item.sale_price", root=row, value=product[amount_key])
        with allure.step(f"27 - {currency}: narxlarni saqlash"):
            user_base.click(name="Сохранить", exact=True)
            user_base.confirm_biruni(expected_text="Сохранить?")
            user_base.expect_page(heading="Цены", url="price_type_list")
        with allure.step(f"28 - {currency}: Цены listini qayta ochish"):
            user_base.navigate_to(tab="Справочники", name="Цены")
            user_base.expect_page(heading="Цены", url="price_type_list")
        with allure.step(f"29 - {currency}: saqlangan narx formasini qayta ochish"):
            user_base.grid_controller(search=price_type)
            user_base.grid(price_type, currency, click=True)
            user_base.click(name="Установка цен", exact=True)
            user_base.expect_page(heading=f"Установка цен: {price_type}", url="product_price?price_type_id=")
        for product in products:
            with allure.step(f"30 - {product['name']}: saqlangan {currency} narxi"):
                user_base.input(ng_model="g.searchValue", root='b-pg-controller[name="products"]', value=product["code"]).press("Enter")
                row = user_base.grid(product["code"], product["name"], root='b-pg-grid[name="products"]')
                user_base.input(ng_model="item.sale_price", root=row, expect_value=product[amount_key])
        with allure.step(f"31 - {currency}: narx formasini saqlamasdan yopish"):
            user_base.click(name="Закрыть", exact=True)
            user_base.confirm_biruni("Сохранить внесенные изменения?", button_name="нет")
            user_base.expect_page(heading="Цены", url="price_type_list")

    with allure.step("32 - Tekshirilgan beshta mahsulotni test yonidagi data_store.json ga yozish"):
        for product in products:
            save_data(f"regression_tmz_{product['number']}", {**product, "run_code": str(code), "server": os.environ["COMPANY_URL"].rstrip("/"), "company": current_company_code(), "card_verified": True, "price_verified": True, "stock_verified": False}, path=PRODUCT_DATA_PATH)


@allure.title("TMZ-2/3/4/5/6: Excel import, view va Цены listidan UZS/USD narxlash")
def test_02_import_tmz(page, group_user_page, code):
    run_import_product(page, group_user_page, code)
